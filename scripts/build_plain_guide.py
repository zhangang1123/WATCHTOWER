#!/usr/bin/env python3
"""Build the plain-language Watchtower-Lite project guide."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "docs" / "Watchtower-Lite项目通俗解读.docx"

# compact_reference_guide preset, with a named CJK font override.
FONT = "Microsoft YaHei"
MONO = "Consolas"
NAVY = RGBColor(24, 56, 96)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
CYAN = RGBColor(8, 145, 178)
INK = RGBColor(35, 43, 58)
MUTED = RGBColor(104, 114, 132)
LIGHT_BLUE = "E8EEF5"
LIGHT_CYAN = "EAF7FA"
LIGHT_GRAY = "F4F6F9"
LIGHT_GREEN = "EAF7F1"
LIGHT_AMBER = "FFF5DF"
LIGHT_RED = "FDECEF"
BORDER = "D6DEE9"
TABLE_WIDTH = 9360
TABLE_INDENT = 120


def set_run_font(run, size=11, bold=False, color=INK, name=FONT, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    return run


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths):
    assert sum(widths) == TABLE_WIDTH
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[min(index, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def add_numbering_definition(doc, fmt, text, font=FONT):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.extend([tabs, ind])
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), font)
    fonts.set(qn("w:hAnsi"), font)
    fonts.set(qn("w:eastAsia"), font)
    r_pr.append(fonts)
    lvl.extend([start, num_fmt, lvl_text, suff, p_pr, r_pr])
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list_item(doc, text, num_id, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)
    if bold_lead and text.startswith(bold_lead):
        set_run_font(p.add_run(bold_lead), bold=True)
        set_run_font(p.add_run(text[len(bold_lead):]))
    else:
        set_run_font(p.add_run(text))
    return p


def add_paragraph(doc, text="", *, bold_lead=None, color=INK, italic=False, after=6, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        set_run_font(p.add_run(bold_lead), bold=True, color=color)
        set_run_font(p.add_run(text[len(bold_lead):]), color=color, italic=italic)
    else:
        set_run_font(p.add_run(text), color=color, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_callout(doc, label, text, fill=LIGHT_CYAN, accent=CYAN):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [TABLE_WIDTH])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.25
    set_run_font(p.add_run(f"{label}  "), bold=True, color=accent)
    set_run_font(p.add_run(text), color=INK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)
    return table


def add_table(doc, headers, rows, widths, header_fill=LIGHT_BLUE, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, header in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(header), size=font_size, bold=True, color=NAVY)
    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            if i == 0 and len(headers) <= 3:
                set_run_font(p.add_run(str(value)), size=font_size, bold=True, color=DARK_BLUE)
            else:
                set_run_font(p.add_run(str(value)), size=font_size, color=INK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_definition(doc, term, simple, detail=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(term), size=11, bold=True, color=DARK_BLUE)
    set_run_font(p.add_run(f"  {simple}"), size=10.5, color=INK)
    if detail:
        p2 = doc.add_paragraph()
        p2.paragraph_format.left_indent = Inches(0.18)
        p2.paragraph_format.space_after = Pt(5)
        p2.paragraph_format.line_spacing = 1.2
        set_run_font(p2.add_run(detail), size=9.5, color=MUTED)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in specs.items():
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_page(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run("WATCHTOWER-LITE  ·  项目通俗解读"), size=8.5, bold=True, color=MUTED)

    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    set_table_geometry(table, [7200, 2160])
    table._tbl.tblPr.remove(table._tbl.tblPr.find(qn("w:tblBorders"))) if table._tbl.tblPr.find(qn("w:tblBorders")) is not None else None
    left = table.cell(0, 0).paragraphs[0]
    set_run_font(left.add_run("从告警到诊断，再到安全处置"), size=8.5, color=MUTED)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(right.add_run("第 "), size=8.5, color=MUTED)
    add_page_field(right)
    set_run_font(right.add_run(" 页"), size=8.5, color=MUTED)


def add_cover(doc):
    for _ in range(4):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    set_run_font(p.add_run("通俗项目指南"), size=11, bold=True, color=CYAN)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    set_run_font(title.add_run("完全理解 Watchtower-Lite"), size=28, bold=True, color=NAVY)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(22)
    set_run_font(subtitle.add_run("一个会接收告警、调查证据、判断根因并尝试处置的智能运维系统"), size=13, color=DARK_BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(70)
    set_run_font(p.add_run("不要求你预先懂 Go、Python、Kubernetes 或 AI Agent"), size=10.5, italic=True, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    set_run_font(p.add_run("WATCHTOWER-LITE"), size=10, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("项目理解手册  ·  2026 年 7 月"), size=9.5, color=MUTED)
    doc.add_page_break()


def build():
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    bullet_id = add_numbering_definition(doc, "bullet", "•", FONT)
    decimal_id = add_numbering_definition(doc, "decimal", "%1.", FONT)
    doc.core_properties.title = "完全理解 Watchtower-Lite"
    doc.core_properties.subject = "Watchtower-Lite 项目通俗解读"
    doc.core_properties.author = "Watchtower-Lite"
    doc.core_properties.keywords = "AIOps, 告警, Incident, Agent, Kubernetes"

    add_cover(doc)

    add_heading(doc, "先用一句话理解它", 1)
    add_callout(doc, "一句话", "Watchtower-Lite 像一名永远在线的初级值班工程师：收到系统告警后，它会整理告警、收集现场证据、判断可能的根因，再根据风险决定是模拟修复还是交给人工。")
    add_paragraph(doc, "它不是一个单纯展示监控图表的页面，也不是一个只会聊天的 AI。它真正想表达的是一条完整的运维处理链：告警进来以后，系统要知道“发生了什么、为什么发生、下一步能不能安全执行”。")
    add_heading(doc, "你可以把它想成一座机场塔台", 2)
    add_paragraph(doc, "机场里同时有雷达、跑道、飞机、调度员和应急预案。Watchtower-Lite 也有类似分工：监控告警像雷达信号，Gateway 像塔台入口，Brain 像分析情况的调度员，自动修复策略像应急预案，Web Console 则是所有人共同看到的指挥屏。")
    add_table(doc, ["现实中的角色", "项目里的角色", "负责什么"], [
        ("雷达", "Prometheus / 告警来源", "发现指标异常并发出告警。"),
        ("塔台入口", "Gateway", "接收、清洗、去重、排队和聚合告警。"),
        ("调度员", "Brain Agent", "查询指标、集群事件和日志，形成诊断。"),
        ("应急预案", "AutoFix Engine", "判断动作风险，决定能否自动执行。"),
        ("指挥大屏", "Web Console", "实时展示事件状态、证据和处理结果。"),
        ("训练场", "Mock Infra", "在没有真实集群时制造可信的演示故障数据。"),
    ], [1800, 2340, 5220])

    add_heading(doc, "为什么需要这个项目", 1)
    add_heading(doc, "传统故障处理有什么痛点", 2)
    add_list_item(doc, "告警很多，但大量告警其实描述的是同一个问题。", bullet_id)
    add_list_item(doc, "值班人员要在监控、Kubernetes、日志系统之间反复切换。", bullet_id)
    add_list_item(doc, "不同工程师经验不同，排查速度和结论质量不稳定。", bullet_id)
    add_list_item(doc, "重复故障每次都从头调查，历史经验没有真正沉淀。", bullet_id)
    add_list_item(doc, "自动修复很危险：如果没有风险控制，错误动作可能扩大故障。", bullet_id)
    add_heading(doc, "Watchtower-Lite 的回答", 2)
    add_paragraph(doc, "项目把一连串人工动作变成了可观察、可控制的程序流程。重点不是“AI 猜得多聪明”，而是让每个判断都有证据，让每个状态都能追踪，让每个自动动作都先经过安全规则。")
    add_callout(doc, "关键理解", "这个项目最有价值的地方是“流程编排和安全边界”，不是当前规则推理本身。真实 LLM 可以以后替换，但告警接入、事件状态机、证据工具和修复护栏仍然需要保留。", LIGHT_AMBER, RGBColor(183, 118, 18))

    doc.add_page_break()
    add_heading(doc, "一次故障是怎样被处理的", 1)
    add_paragraph(doc, "下面用 payment 服务内存耗尽（OOMKilled）作为例子。你在前端点击“注入故障”后，真实发生的是以下十步。")
    steps = [
        ("创建演示现场。", "Gateway 先通知 Mock Infra：payment 服务现在进入 oom_kill 场景。Mock Infra 随即准备异常指标、OOM 事件和内核日志。"),
        ("产生一条告警。", "系统创建 P0 级 OOMKilled 告警。P0 表示最高优先级，需要立即处理。"),
        ("标准化。", "不同来源的字段被转换成项目内部统一的 Alert 格式。后面的流程不用再关心告警来自哪里。"),
        ("去重和静默。", "系统计算稳定指纹；五分钟内完全相同的告警不会重复处理。符合静默规则的告警也会被拦截。"),
        ("进入有界优先级队列。", "P0 会排在 P1、P2、P3 前面；队列有容量上限，避免告警风暴无限占用内存。"),
        ("聚合为 Incident。", "同一服务、同一故障类别的相关告警会合并成一个运维事件，而不是制造多张重复工单。"),
        ("Brain 开始调查。", "Gateway 通过 gRPC 流式调用 Python Brain。Brain 依次查询 Prometheus、K8s 事件和 Pod 日志。"),
        ("形成诊断。", "Brain 根据告警描述和证据形成根因、置信度、证据列表、修复建议和机器可读的 fix_type。"),
        ("风险判断。", "AutoFix Engine 不直接执行自然语言，而是检查 fix_type、风险等级和置信度阈值。高风险动作会升级人工。"),
        ("实时展示。", "每次状态变化和调查步骤通过 WebSocket 推送到控制台，页面无需反复刷新。"),
    ]
    for title, detail in steps:
        add_list_item(doc, f"{title} {detail}", decimal_id, bold_lead=title)

    add_heading(doc, "为什么要把 Alert 和 Incident 分开", 2)
    add_definition(doc, "Alert（告警）", "一个监控信号。", "例如 CPU 超过 90%、Pod 被 OOMKilled。它描述的是一个现象，而且可能短时间内重复出现很多次。")
    add_definition(doc, "Incident（事件）", "一个需要跟踪到结束的故障任务。", "它可以包含多条相关告警，并拥有状态、诊断、修复过程和恢复时间。")
    add_callout(doc, "类比", "Alert 像多位乘客同时报警“闻到烟味”；Incident 是消防部门建立的一次正式火警任务。几十通电话不应该变成几十场独立火灾。")

    doc.add_page_break()
    add_heading(doc, "系统由哪些部分组成", 1)
    add_table(doc, ["服务", "默认端口", "通俗解释", "主要技术"], [
        ("Web Console", "3000（开发）", "浏览器里的指挥大屏。", "Vue 3 + Vite"),
        ("Gateway", "8080", "所有告警和前端请求的统一入口。", "Go + Gin"),
        ("Brain", "50052", "负责调查和生成诊断的“大脑”。", "Python + gRPC"),
        ("Mock K8s", "8081", "提供模拟集群事件、Pod 和日志。", "Go + Gin"),
        ("Mock Prometheus", "9090", "提供模拟监控指标查询。", "Go + Gin"),
    ], [1800, 1200, 3900, 2460])

    add_heading(doc, "Gateway：告警入口和流程总管", 2)
    add_paragraph(doc, "Gateway 使用 Go 编写。Go 很适合处理大量并发网络请求，因此它负责接收 webhook、校验输入、去重、排队、创建 Incident、维护状态和连接前端。")
    add_list_item(doc, "Normalizer：把不同来源的告警翻译成统一格式。", bullet_id, "Normalizer：")
    add_list_item(doc, "Deduplicator：在时间窗口内识别重复告警。", bullet_id, "Deduplicator：")
    add_list_item(doc, "Silencer：根据规则忽略维护期或已知噪声。", bullet_id, "Silencer：")
    add_list_item(doc, "Priority Queue：让最严重的告警优先处理，同时限制最大容量。", bullet_id, "Priority Queue：")
    add_list_item(doc, "Incident Manager：聚合告警并推动事件状态变化。", bullet_id, "Incident Manager：")

    add_heading(doc, "Brain：调查证据并生成诊断", 2)
    add_paragraph(doc, "Brain 使用 Python 编写，因为 Python 的 AI、数据处理和工具生态更丰富。当前版本虽然保留了 LLM 配置，但主要推理仍然是确定性规则：先查指标，再查 K8s 事件，再查日志，最后按故障关键词形成结论。")
    add_callout(doc, "不要误解", "当前 Brain 不是一个自由规划、自由调用工具的真实大模型 Agent。它更像一个可预测的 Agent 框架样板。这样做方便离线演示，也让测试结果稳定。", LIGHT_RED, RGBColor(181, 52, 72))

    add_heading(doc, "Mock Infra：可以控制的故障训练场", 2)
    add_paragraph(doc, "真实 Prometheus 和 Kubernetes 环境很重，也不适合为了演示故意制造故障。Mock Infra 用普通 HTTP 接口模拟它们。关键在于三类证据必须相互印证：指标显示异常，K8s 事件描述异常，日志也出现对应错误。")

    add_heading(doc, "Web Console：把内部过程变成可见任务", 2)
    add_paragraph(doc, "控制台不是简单的监控图表。它围绕 Incident 展示优先级、当前状态、根因、置信度、调查轨迹、修复建议和实时事件总线。明暗模式偏好会保存到当前浏览器。")

    doc.add_page_break()
    add_heading(doc, "事件为什么需要状态机", 1)
    add_paragraph(doc, "如果没有明确状态，系统可能在尚未诊断时就执行修复，或者已经解决后又回到诊断中。状态机就是一套“只允许合法步骤”的交通规则。")
    add_table(doc, ["状态", "含义", "通常下一步"], [
        ("new", "事件刚创建。", "diagnosing 或 false_alarm"),
        ("diagnosing", "Brain 正在收集证据。", "diagnosed 或 escalated"),
        ("diagnosed", "已经形成根因和建议。", "fixing、resolved 或 escalated"),
        ("fixing", "正在执行模拟修复。", "fixed 或 escalated"),
        ("fixed", "动作完成，进入观察。", "resolved 或 escalated"),
        ("resolved", "确认恢复，事件结束。", "终态"),
        ("escalated", "自动化不适合继续，交给人工。", "人工处理后 resolved"),
        ("false_alarm", "确认是误报。", "终态"),
    ], [1700, 4300, 3360])
    add_paragraph(doc, "当前 Manager 会通过统一锁保护 Incident 的读写，再返回快照给 API 和 WebSocket。这样诊断 goroutine 修改状态时，前端查询不会读到一半更新的数据。")

    add_heading(doc, "系统之间如何通信", 1)
    add_heading(doc, "HTTP / REST", 2)
    add_paragraph(doc, "浏览器查询 Incident、触发模拟故障，以及 Brain 查询 Mock 指标和日志时使用 HTTP。它适合一问一答。")
    add_heading(doc, "gRPC", 2)
    add_paragraph(doc, "Gateway 调用 Brain 时使用 gRPC。gRPC 是一种服务间远程调用协议：双方先用 .proto 文件约定字段和方法，再生成 Go、Python 代码。这样比手写 JSON 更严格。")
    add_heading(doc, "流式 gRPC", 2)
    add_paragraph(doc, "普通调用只在最后返回一次结果；流式调用可以一边调查、一边返回每一步。Watchtower-Lite 用它把“查询指标→查询事件→查询日志”的过程逐步传回 Gateway。")
    add_heading(doc, "WebSocket", 2)
    add_paragraph(doc, "WebSocket 是浏览器与服务器之间的长连接。服务器不必等浏览器来问，就能主动推送新 Incident、状态变化、诊断步骤和修复结果。项目还加入了 ping/pong 心跳和读写超时，避免死连接长期占用资源。")

    doc.add_page_break()
    add_heading(doc, "自动修复为什么不能直接相信 AI", 1)
    add_paragraph(doc, "自然语言很适合给人阅读，却不适合直接执行。例如“重启一下服务”没有说明集群、命名空间、工作负载和最大影响范围。项目因此把展示文本与机器动作分开。")
    add_definition(doc, "suggested_fix", "给人看的修复建议。", "可以是多行自然语言，解释临时方案和长期方案。")
    add_definition(doc, "fix_type", "给程序看的动作类型。", "例如 restart_pod、scale、cleanup_disk 或 config_change。AutoFix Engine 只根据结构化动作做风险判断。")
    add_definition(doc, "confidence（置信度）", "系统对诊断结论把握有多大。", "通常是 0 到 1，例如 0.92 代表 92%。它不是事实证明，只是风险决策中的一个输入。")
    add_table(doc, ["动作类型", "当前风险判断", "系统行为"], [
        ("restart_pod", "低风险，置信度至少 85%", "可以进入模拟自动修复。"),
        ("scale", "低风险，置信度至少 90%", "满足阈值后模拟扩容。"),
        ("cleanup_disk", "中风险", "不自动执行，等待人工确认。"),
        ("config_change", "高风险", "直接升级人工。"),
        ("traffic_switch", "高风险", "直接升级人工。"),
    ], [2200, 3200, 3960])
    add_callout(doc, "安全底线", "未来接入真实 Kubernetes 后，也不能把 LLM 生成的文字拼进 shell 命令。正确方式是使用受限、类型化、可审计的动作接口，并设置作用范围、审批、验证和回滚。", LIGHT_AMBER, RGBColor(183, 118, 18))

    add_heading(doc, "项目支持哪些演示故障", 1)
    add_table(doc, ["场景", "你会看到什么", "典型结论"], [
        ("oom_kill", "服务不可用、OOMKilling 事件、内核杀进程日志", "内存限制不足或内存使用过高"),
        ("high_cpu", "CPU 使用率 95%、CPU throttling", "计算资源不足，需要扩容或优化"),
        ("high_latency", "P99 延迟 2.5 秒、数据库连接池耗尽日志", "数据库连接问题导致变慢"),
        ("disk_full", "磁盘使用率 92%、DiskPressure", "日志堆积或容量不足"),
        ("conn_timeout", "高错误率、Redis 连接超时日志", "下游连接数或连接池问题"),
        ("crash_loop", "Pod 不断重启、panic 和越界日志", "代码异常导致 CrashLoop"),
    ], [1700, 4380, 3280])

    doc.add_page_break()
    add_heading(doc, "怎样运行和观察项目", 1)
    add_heading(doc, "启动", 2)
    add_paragraph(doc, "在项目根目录使用 Git Bash 运行 ./scripts/start.sh。脚本会检查依赖和端口、编译 Go 服务，然后把真实进程交给 Python Supervisor 管理。")
    add_table(doc, ["地址", "用途"], [
        ("http://localhost:3000", "开发模式下的 Web Console"),
        ("http://localhost:8080", "Gateway API 与生产静态页面入口"),
        ("ws://localhost:8080/ws", "WebSocket 实时通道"),
        ("127.0.0.1:50052", "Brain gRPC 服务"),
        ("http://localhost:9090", "Mock Prometheus"),
        ("http://localhost:8081", "Mock K8s 与场景控制"),
    ], [2700, 6660])
    add_heading(doc, "为什么 Ctrl+C 现在能可靠关闭", 2)
    add_paragraph(doc, "旧脚本在 Git Bash 中记录到的可能是 MSYS 中间 PID，而不是真实 node.exe PID，因此父进程退出后 Node 还占着 3000。现在 Supervisor 直接通过 Python subprocess 持有 Windows 真实 PID，Ctrl+C、SIGTERM 或任一服务异常退出都会按进程树清理全部服务。")
    add_heading(doc, "端口冲突怎么判断", 2)
    add_list_item(doc, "本地脚本和 Docker Compose 使用同一组宿主机端口，不能同时运行。", bullet_id)
    add_list_item(doc, "Vite 使用 strictPort，3000 被占用时直接报错，不会悄悄改成 3001。", bullet_id)
    add_list_item(doc, "启动前会一次性检查 3000、8080、8081、9090 和 50052。", bullet_id)
    add_list_item(doc, "可以运行 ./scripts/smoke-test.sh 验证完整启动、诊断和端口释放。", bullet_id)

    add_heading(doc, "前端应该怎么看", 1)
    add_list_item(doc, "先看左侧服务状态，确认 Gateway、Brain 和 Mock Infra 正常。", decimal_id)
    add_list_item(doc, "点击“注入故障”，选择故障场景和目标服务。", decimal_id)
    add_list_item(doc, "观察事件队列里出现新的 P0/P1/P2 任务。", decimal_id)
    add_list_item(doc, "打开事件详情，看状态从 new 进入 diagnosing。", decimal_id)
    add_list_item(doc, "查看调查轨迹，理解 Brain 查询了哪些证据。", decimal_id)
    add_list_item(doc, "查看根因、置信度、fix_type 和建议动作。", decimal_id)
    add_list_item(doc, "在实时事件总线观察 WebSocket 推送。", decimal_id)

    doc.add_page_break()
    add_heading(doc, "哪些能力是真的，哪些仍是模拟的", 1)
    add_paragraph(doc, "这是理解项目边界最重要的一节。Watchtower-Lite 已经具备完整的软件骨架，但它还不是可以直接接管生产集群的 AIOps 产品。")
    add_table(doc, ["能力", "当前真实程度", "说明"], [
        ("告警 HTTP 接入", "真实", "可以接收 Prometheus 和自定义 webhook。"),
        ("去重、排队、聚合", "真实（内存版）", "逻辑真实，但服务重启后数据会丢失。"),
        ("Incident 状态机", "真实", "非法状态转换会被拒绝。"),
        ("gRPC 流式诊断", "真实", "Go 与 Python 会实际进行流式通信。"),
        ("WebSocket 实时推送", "真实", "前端会接收真实状态事件。"),
        ("Prometheus / K8s 数据", "模拟", "来自 Mock Infra，不是真实集群。"),
        ("AI 推理", "规则驱动", "框架像 Agent，但当前结论主要由关键词规则产生。"),
        ("历史向量检索", "简化模拟", "实际是内存关键词匹配，不是真实 embedding。"),
        ("自动修复执行", "模拟", "只记录动作日志，不会修改真实 Kubernetes。"),
        ("修复后验证", "简化模拟", "等待观察期后假设成功，尚未查询真实恢复指标。"),
        ("持久化", "未实现", "Incident 和历史模式都在内存中。"),
        ("认证与权限", "未实现", "当前适合本地演示，不应直接暴露到公网。"),
    ], [2200, 1900, 5260])
    add_callout(doc, "正确定位", "它已经是一套可运行、可演示、可继续工程化的 AIOps 原型；它还不是一个经过安全审计、拥有真实执行权限和生产数据保障的平台。", LIGHT_GREEN, RGBColor(19, 128, 91))

    add_heading(doc, "关键术语词典", 1)
    definitions = [
        ("AIOps", "用数据和自动化帮助运维。", "AI for IT Operations。它不等于“让 AI 随便操作服务器”，而是分析告警、关联事件、辅助诊断和自动化处置。"),
        ("Webhook", "别人有事时主动通知你的 HTTP 地址。", "Alertmanager 触发告警后，会向 Gateway 的 webhook 发送 JSON。"),
        ("Prometheus", "专门存储和查询监控指标的系统。", "例如 CPU 使用率、错误率、延迟和服务是否在线。"),
        ("PromQL", "Prometheus 的查询语言。", "类似数据库的 SQL，但主要用于查询时间序列指标。"),
        ("Kubernetes / K8s", "管理容器化应用的集群平台。", "它负责部署、扩缩容、重启和调度服务。K8s 是 Kubernetes 的常用缩写。"),
        ("Pod", "Kubernetes 中运行容器的基本单元。", "可以把它粗略理解为一个服务实例。"),
        ("OOMKilled", "进程因为内存不足被系统杀死。", "OOM 是 Out Of Memory。它可能由内存限制太低、内存泄漏或流量突增引起。"),
        ("CrashLoopBackOff", "容器启动后反复崩溃，K8s 暂缓重启。", "它是现象，不直接等于根因；真正原因仍要查日志和事件。"),
        ("gRPC", "服务之间的强类型远程调用方式。", "通过 protobuf 约定接口，适合 Go、Python 等不同语言之间通信。"),
        ("Protocol Buffers / protobuf", "一种接口和数据结构说明书。", ".proto 文件像双方共同签署的表格模板，字段变化后要重新生成代码。"),
        ("WebSocket", "浏览器和服务器之间保持不断开的双向连接。", "适合实时推送，而不是每几秒重新请求一次。"),
        ("Goroutine", "Go 中轻量的并发任务。", "可以同时处理多个告警、诊断和观察任务，但共享数据必须避免竞态。"),
        ("Race Condition / 数据竞争", "多个并发任务同时读写同一数据导致结果不确定。", "项目通过 Manager 锁和快照减少这类问题。"),
        ("State Machine / 状态机", "规定对象允许怎样变化的一套规则。", "例如 diagnosed 可以进入 fixing，但 resolved 不应该重新回到 diagnosing。"),
        ("Backpressure / 背压", "下游处理不过来时，上游不能无限灌数据。", "有界队列就是一种基础背压机制。"),
        ("MTTR", "平均修复时间。", "Mean Time To Repair/Resolve，从故障发生到恢复所花的时间。"),
        ("Agent", "能根据目标调用工具并推进任务的软件角色。", "真正 Agent 应根据上下文选择下一步；当前项目主要使用固定调查顺序。"),
        ("ReAct", "边推理、边行动、边观察的 Agent 模式。", "名称来自 Reasoning + Acting。项目 UI 展示调查步骤，但不应把模型私有思维链当作产品功能。"),
        ("LLM", "大语言模型。", "例如用于理解日志、选择工具和生成结论。当前项目尚未真正依赖 LLM。"),
        ("RAG", "先检索资料，再让模型回答。", "Retrieval-Augmented Generation。未来可用来检索 Runbook、历史事故和服务文档。"),
        ("Runbook", "处理某类故障的操作手册。", "它规定检查什么、怎样修复、怎样验证和怎样回滚。"),
        ("RBAC", "按角色控制权限。", "Role-Based Access Control。例如查看者只能看，审批者可以批准修复。"),
    ]
    for term, simple, detail in definitions:
        add_definition(doc, term, simple, detail)

    doc.add_page_break()
    add_heading(doc, "常见疑问", 1)
    qa = [
        ("没有 OpenAI API Key 能运行吗？", "能。当前 Brain 使用规则推理，演示链路不依赖 OpenAI API。"),
        ("它会真的修改 Kubernetes 吗？", "不会。当前 AutoFix Executor 只打印模拟动作；控制台也明确显示 Simulation Mode。"),
        ("为什么既用 Go 又用 Python？", "Go 负责高并发入口和可靠状态流程；Python 负责 Agent、工具和未来 AI 生态。两者通过 gRPC 解耦。"),
        ("为什么不用一个服务全部写完？", "小项目可以，但分离后更容易独立扩展和替换 Brain，也能展示跨语言服务设计。代价是部署和协议维护更复杂。"),
        ("服务重启后历史事件还在吗？", "不在。当前存储是内存 map。要长期使用，需要接入 SQLite 或 PostgreSQL。"),
        ("为什么 P2/P3 不马上诊断？", "它们先观察十分钟，用来减少低优先级噪声。实际生产中应把观察时间做成配置。"),
        ("为什么诊断后经常升级人工？", "这是安全设计。结构化动作不在低风险白名单，或置信度未达到阈值时，不应该自动执行。"),
        ("它能横向扩容多个 Gateway 吗？", "当前不适合。去重、队列、Incident 和状态都在单进程内存中。横向扩容需要共享存储和消息系统。"),
        ("能直接放到公网吗？", "不能。当前缺少完整认证、RBAC、Webhook 签名、TLS、审计和限流。"),
    ]
    for question, answer in qa:
        add_definition(doc, question, answer)

    add_heading(doc, "如果继续开发，应该先做什么", 1)
    roadmap = [
        ("第一阶段：可靠的单机 MVP", "SQLite/PostgreSQL 持久化、分页、配置管理、更多测试、真实健康验证。"),
        ("第二阶段：安全接入真实环境", "只读 Prometheus/Kubernetes、API 鉴权、RBAC、审计日志、密钥脱敏。"),
        ("第三阶段：真实 Agent", "接入 LLM、严格 JSON Schema、可选择工具、RAG 检索 Runbook 和历史事件。"),
        ("第四阶段：受控自动修复", "Dry-run、人工审批、Canary、影响范围限制、真实验证、回滚和 Kill Switch。"),
        ("第五阶段：多实例平台化", "事件总线、分布式锁、共享去重、多租户隔离、SLO 和成本观测。"),
    ]
    add_table(doc, ["阶段", "核心目标"], roadmap, [2700, 6660])

    add_heading(doc, "读代码时从哪里开始", 1)
    add_table(doc, ["想理解的内容", "先看文件"], [
        ("整个系统怎样启动", "scripts/start.sh、scripts/supervisor.py"),
        ("HTTP 路由和服务装配", "cmd/gateway/main.go"),
        ("告警接入和队列", "internal/gateway/handler.go、queue.go"),
        ("告警与事件模型", "internal/models/alert.go、incident.go、diagnosis.go"),
        ("事件状态流转", "internal/incident/manager.go、state_machine.go"),
        ("Brain 调用", "internal/agent/client.go、proto/diagnosis.proto"),
        ("诊断规则和工具", "brain/agent.py、brain/tools/"),
        ("自动修复策略", "internal/autofix/engine.go"),
        ("实时事件", "internal/ws/hub.go、web/src/api.js"),
        ("前端页面", "web/src/App.vue、web/src/style.css"),
    ], [3000, 6360])

    doc.add_page_break()
    add_heading(doc, "现在你应该能这样介绍这个项目", 1)
    add_callout(doc, "60 秒版本", "Watchtower-Lite 是一个 Go 与 Python 混合架构的 AIOps 原型。Go Gateway 负责接收、去重、排队和聚合告警，并通过状态机管理 Incident；Python Brain 通过 gRPC 流式查询模拟 Prometheus、Kubernetes 事件和日志，形成带证据和置信度的诊断；AutoFix Engine 根据结构化 fix_type 和风险阈值决定模拟修复或升级人工；Vue 控制台通过 WebSocket 实时展示全过程。当前基础流程是真实可运行的，但监控数据、AI 推理和修复执行仍以规则与 Mock 为主，下一步需要持久化、安全权限和真实基础设施接入。", LIGHT_GREEN, RGBColor(19, 128, 91))
    add_heading(doc, "最后记住三件事", 2)
    add_list_item(doc, "Alert 是信号，Incident 是需要跟踪到底的故障任务。", decimal_id)
    add_list_item(doc, "Agent 的价值不只是给答案，而是通过工具形成可验证的证据链。", decimal_id)
    add_list_item(doc, "自动修复必须有结构化动作、风险护栏、验证和回滚，不能直接执行 AI 文本。", decimal_id)
    add_paragraph(doc, "理解了这三点，你就已经抓住了 Watchtower-Lite 的核心设计。", bold_lead="理解了这三点，", after=0)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
